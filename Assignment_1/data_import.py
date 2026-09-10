"""
Part 0 Step 2 data build — run this once to produce your locked dataset.

Reads tickers from Holdings!B3:B52 in portfolio.xlsx, downloads adjusted
monthly closes for those 50 tickers + ^GSPC via yfinance, aligns to a
complete-case date range, drops any ticker with >5% missing months,
converts to simple returns, and writes:
  - a "Prices" sheet with monthly adjusted close prices
  - a "Returns" sheet with monthly simple returns + a red/white/green
    color scale
  - a "Data Notes" sheet documenting anything dropped or flagged
Formatting (fonts, colors, header style, row banding) matches the
existing Holdings/Overview sheets. Also writes two CSVs. Overwrites
the workbook in place.
"""

import openpyxl
import pandas as pd
import yfinance as yf
from datetime import datetime
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill

# ---- EDIT THIS to the path of your file on your own machine ----
SRC = "portfolio.xlsx"
OUT = "portfolio.xlsx"  # overwrite in place

START = "2010-01-01"
END = datetime.today().strftime("%Y-%m-%d")

# Shared style constants, matched to the existing Holdings sheet
FONT_NAME = "Arial"
TITLE_FILL = PatternFill(fill_type="solid", fgColor="FF17324A")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="FF24476A")
BAND_FILL_LIGHT = PatternFill(fill_type="solid", fgColor="FFF2F5F8")
BAND_FILL_WHITE = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
TITLE_FONT = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFFFF")
HEADER_FONT = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=12, bold=False, color="FF222222")
DATE_FONT = Font(name=FONT_NAME, size=12, bold=True, color="FF17324A")

# ---------------------------------------------------------------
# 1. Read tickers from Holdings (col B, rows 3-52)
# ---------------------------------------------------------------
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Holdings"]
tickers = []
for row in range(3, 53):
    t = ws.cell(row=row, column=2).value  # column B = Ticker
    if t:
        tickers.append(str(t).strip().upper())
tickers = sorted(set(tickers))
print(f"Read {len(tickers)} tickers")

MARKET = "^GSPC"
all_symbols = tickers + [MARKET]

# ---------------------------------------------------------------
# 2. Download adjusted monthly closes
# ---------------------------------------------------------------
print("Downloading monthly data...")
raw = yf.download(
    all_symbols, start=START, end=END, interval="1mo",
    auto_adjust=True, progress=False, group_by="ticker"
)

prices = pd.DataFrame()
missing_symbols = []
for sym in all_symbols:
    try:
        col = raw[sym]["Close"]
        if col.dropna().empty:
            missing_symbols.append(sym)
            continue
        prices[sym] = col
    except KeyError:
        missing_symbols.append(sym)

if missing_symbols:
    print(f"\n*** DROPPED — no data returned at all: {missing_symbols} ***")

prices.index = pd.to_datetime(prices.index)
print(f"Raw panel shape: {prices.shape}")

# ---------------------------------------------------------------
# 3. Drop tickers with >5% missing monthly observations
# ---------------------------------------------------------------
missing_frac = prices.isna().mean()
dropped_for_missing = missing_frac[missing_frac > 0.05].index.tolist()
dropped_missing_pct = {t: round(missing_frac[t] * 100, 1) for t in dropped_for_missing}
if dropped_for_missing:
    print(f"\n*** DROPPED — >5% missing monthly observations: ***")
    for t, pct in dropped_missing_pct.items():
        print(f"    {t}: {pct}% missing")
else:
    print("\nNo tickers dropped for missing data (all under 5%).")
prices = prices.drop(columns=dropped_for_missing)

# ---------------------------------------------------------------
# 4. Complete-case alignment (keep only months every symbol has data)
# ---------------------------------------------------------------
before_rows = len(prices)
prices = prices.dropna(how="any")
after_rows = len(prices)
rows_dropped_for_alignment = before_rows - after_rows
print(f"\nComplete-case alignment: {before_rows} -> {after_rows} months "
      f"({rows_dropped_for_alignment} month(s) dropped)")
print(f"Sample period: {prices.index.min().date()} to {prices.index.max().date()}")
print(f"Final ticker count (excl. market): {prices.shape[1]-1}")

# ---------------------------------------------------------------
# 5. Convert to simple returns (after cleaning)
# ---------------------------------------------------------------
returns = prices.pct_change().dropna(how="any")

# Flag extreme returns (< -50% or > +200%) -- flag only, never auto-remove
extreme = returns[(returns < -0.50) | (returns > 2.00)]
extreme_flags = []
for c in extreme.columns:
    for dt, val in extreme[c].dropna().items():
        extreme_flags.append((dt.date(), c, val))
if extreme_flags:
    print(f"\n*** FLAGGED (not removed) — {len(extreme_flags)} extreme monthly return(s): ***")
    for dt, c, val in extreme_flags:
        print(f"    {dt} {c} {val:.2%}")
else:
    print("\nNo extreme monthly returns flagged.")

prices.to_csv("monthly_prices.csv")
returns.to_csv("monthly_returns.csv")

# ---------------------------------------------------------------
# 6. Write into "Prices", "Returns", and "Data Notes" sheets
# ---------------------------------------------------------------
wb2 = openpyxl.load_workbook(SRC, data_only=False)


def write_sheet(wb, sheet_name, df, number_format, is_returns=False):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    sh = wb.create_sheet(sheet_name)

    # Title row (row 1), matching the navy title bar on Holdings
    sh.cell(row=1, column=1, value=sheet_name)
    title_cell = sh.cell(row=1, column=1)
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center")
    n_cols_total = len(df.columns) + 1
    for j in range(1, n_cols_total + 1):
        sh.cell(row=1, column=j).fill = TITLE_FILL
    sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols_total)

    # Header row (row 2): Date + tickers
    sh.cell(row=2, column=1, value="Date")
    for j, col in enumerate(df.columns, start=2):
        sh.cell(row=2, column=j, value=col)
    for j in range(1, n_cols_total + 1):
        cell = sh.cell(row=2, column=j)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # Data rows (row 3 onward), alternating band fill to match Holdings
    for i, (dt, rowdata) in enumerate(df.iterrows(), start=3):
        band_fill = BAND_FILL_LIGHT if (i % 2 == 1) else BAND_FILL_WHITE

        date_cell = sh.cell(row=i, column=1, value=dt.date())
        date_cell.font = DATE_FONT
        date_cell.fill = band_fill
        date_cell.number_format = "yyyy-mm-dd"
        date_cell.alignment = Alignment(horizontal="center")

        for j, col in enumerate(df.columns, start=2):
            cell = sh.cell(row=i, column=j, value=float(rowdata[col]))
            cell.font = BODY_FONT
            cell.fill = band_fill
            cell.number_format = number_format
            cell.alignment = Alignment(horizontal="right")

    n_rows_total = len(df) + 2  # + title + header

    # Freeze panes below header row and right of date column
    sh.freeze_panes = "B3"

    # Column widths
    sh.column_dimensions["A"].width = 13
    for j in range(2, n_cols_total + 1):
        sh.column_dimensions[get_column_letter(j)].width = 11

    # Conditional formatting (returns sheet only): red -> white -> green
    if is_returns:
        first_col = get_column_letter(2)
        last_col = get_column_letter(n_cols_total)
        data_range = f"{first_col}3:{last_col}{n_rows_total}"
        rule = ColorScaleRule(
            start_type="min", start_color="FFC0392B",
            mid_type="num", mid_value=0, mid_color="FFFFFFFF",
            end_type="max", end_color="FF1E7B45",
        )
        sh.conditional_formatting.add(data_range, rule)

    return sh


def write_notes_sheet(wb, missing_symbols, dropped_missing_pct, rows_dropped,
                       final_start, final_end, final_n_tickers, extreme_flags):
    sheet_name = "Data Notes"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    sh = wb.create_sheet(sheet_name)

    sh.cell(row=1, column=1, value="Data Build Notes")
    sh.cell(row=1, column=1).font = TITLE_FONT
    sh.cell(row=1, column=1).fill = TITLE_FILL
    for j in range(1, 4):
        sh.cell(row=1, column=j).fill = TITLE_FILL
    sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    sh.cell(row=2, column=1, value=f"Generated: {datetime.today().strftime('%Y-%m-%d %H:%M')}")
    sh.cell(row=2, column=1).font = Font(name=FONT_NAME, size=10, italic=True, color="FF666666")

    r = 4
    labels = [
        ("Final sample period:", f"{final_start} to {final_end}"),
        ("Final ticker count (excl. market):", final_n_tickers),
        ("Months dropped in complete-case alignment:", rows_dropped),
    ]
    for label, value in labels:
        sh.cell(row=r, column=1, value=label).font = Font(name=FONT_NAME, size=12, bold=True, color="FF17324A")
        sh.cell(row=r, column=2, value=value).font = BODY_FONT
        r += 1
    r += 1

    sh.cell(row=r, column=1, value="Tickers dropped — no data returned at all:").font = Font(
        name=FONT_NAME, size=12, bold=True, color="FF17324A")
    r += 1
    if missing_symbols:
        for t in missing_symbols:
            sh.cell(row=r, column=1, value=t).font = BODY_FONT
            r += 1
    else:
        sh.cell(row=r, column=1, value="(none)").font = BODY_FONT
        r += 1
    r += 1

    sh.cell(row=r, column=1, value="Tickers dropped — >5% missing monthly observations:").font = Font(
        name=FONT_NAME, size=12, bold=True, color="FF17324A")
    r += 1
    if dropped_missing_pct:
        sh.cell(row=r, column=1, value="Ticker").font = HEADER_FONT
        sh.cell(row=r, column=1).fill = HEADER_FILL
        sh.cell(row=r, column=2, value="% Missing").font = HEADER_FONT
        sh.cell(row=r, column=2).fill = HEADER_FILL
        r += 1
        for t, pct in dropped_missing_pct.items():
            sh.cell(row=r, column=1, value=t).font = BODY_FONT
            sh.cell(row=r, column=2, value=f"{pct}%").font = BODY_FONT
            r += 1
    else:
        sh.cell(row=r, column=1, value="(none)").font = BODY_FONT
        r += 1
    r += 1

    sh.cell(row=r, column=1,
            value="Extreme monthly returns flagged (< -50% or > +200%, NOT removed):").font = Font(
        name=FONT_NAME, size=12, bold=True, color="FF17324A")
    r += 1
    if extreme_flags:
        for label, col_idx in [("Date", 1), ("Ticker", 2), ("Return", 3)]:
            c = sh.cell(row=r, column=col_idx, value=label)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
        r += 1
        for dt, t, val in extreme_flags:
            sh.cell(row=r, column=1, value=str(dt)).font = BODY_FONT
            sh.cell(row=r, column=2, value=t).font = BODY_FONT
            sh.cell(row=r, column=3, value=f"{val:.2%}").font = BODY_FONT
            r += 1
    else:
        sh.cell(row=r, column=1, value="(none)").font = BODY_FONT

    sh.column_dimensions["A"].width = 48
    sh.column_dimensions["B"].width = 20
    sh.column_dimensions["C"].width = 15

    return sh


write_sheet(wb2, "Prices", prices, number_format="0.0000")
write_sheet(wb2, "Returns", returns, number_format="0.0000%", is_returns=True)
write_notes_sheet(
    wb2, missing_symbols, dropped_missing_pct, rows_dropped_for_alignment,
    prices.index.min().date(), prices.index.max().date(), prices.shape[1] - 1,
    extreme_flags
)

wb2.save(OUT)
print(f"\nSaved -> {OUT}")
print("Wrote 'Prices', 'Returns', and 'Data Notes' sheets, styled to match Holdings.")